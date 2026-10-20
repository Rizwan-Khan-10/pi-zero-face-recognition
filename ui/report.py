"""
Report screen.

Top bar (shared by every tab):
    Branch / Semester / Subject picker  +  "Minimum attendance %" box.
    The minimum starts at config.ATTENDANCE_ELIGIBILITY_THRESHOLD (75) and can be
    typed over at any time - every tab re-colours itself live (red = below minimum).

Tabs:
    Daily            - who was Present / Absent on one date
    Weekly           - one ISO week: attendance by day + per-student %
    Monthly          - one month: daily trend + per-student %
    Semester Summary - the whole semester at a glance: summary cards, charts and a
                       student-wise table with one % column per subject.
"""

from datetime import datetime, timedelta

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import ttkbootstrap as tb
import openpyxl
from openpyxl.styles import Font, PatternFill

import matplotlib
matplotlib.use("TkAgg")
from matplotlib.figure import Figure
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

import config
from core import storage, reporting
from ui.widgets import make_header, make_table

ALL_SUBJECTS_LABEL = "All Subjects"
GREEN, RED, GREY = "#1a7431", "#b3261e", "#888888"
BOLD = ("Segoe UI", 11, "bold")


# ---------------------------------------------------------------- helpers

def _week_key(date_str):
    d = datetime.strptime(date_str, "%Y-%m-%d").date()
    y, w, _ = d.isocalendar()
    return f"{y}-W{w:02d}"


def _month_key(date_str):
    return date_str[:7]


def _week_label(week_key):
    """'2026-W42' -> '12 Oct - 18 Oct 2026' (Monday to Sunday) - easier to read than the ISO week number."""
    monday = datetime.strptime(f"{week_key}-1", "%G-W%V-%u")
    sunday = monday + timedelta(days=6)
    return f"{monday:%d %b} - {sunday:%d %b %Y}"


def _month_label(month_key):
    """'2026-10' -> 'October 2026'."""
    return datetime.strptime(month_key, "%Y-%m").strftime("%B %Y")


def _fmt_pct(percent):
    return "-" if percent is None else f"{percent:.0f}%"


def _subject_cell(present, total, threshold):
    """'82%' or '58% ▼' (below minimum) or '-' (no lectures recorded)."""
    if not total:
        return "-"
    percent = present / total * 100
    return f"{percent:.0f}%" + (" ▼" if percent < threshold else "")


def _empty_chart(fig, ax, canvas, message):
    ax.clear()
    ax.text(0.5, 0.5, message, ha="center", va="center", fontsize=11, color=GREY,
            transform=ax.transAxes)
    ax.set_xticks([])
    ax.set_yticks([])
    for spine in ax.spines.values():
        spine.set_visible(False)
    fig.tight_layout()
    canvas.draw()


def _clear_tree(tree):
    for item in tree.get_children():
        tree.delete(item)


def _fill_simple_tree(tree, rows, threshold):
    """Weekly / Monthly table: Roll, Name, Present, Total, %, Status - red when below minimum."""
    _clear_tree(tree)
    for r in rows:
        low = reporting.is_low(r["pct"], threshold)
        tree.insert("", "end", values=(
            r["roll"], r["name"], r["present"], r["total"],
            _fmt_pct(r["pct"]), reporting.status_text(r["pct"], threshold)),
            tags=("low",) if low else ())


def _save_report_excel(save_path, title, rows, subject_ids, threshold, heading):
    """Excel export used by Weekly / Monthly / Semester. Rows below the minimum are filled red."""
    subjects = storage.load_subjects()
    show_subjects = len(subject_ids) > 1
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (title or "Report")[:31]

    ws.append([f"{heading}   |   minimum attendance: {threshold:g}%"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])

    header = ["Roll No", "Name"]
    if show_subjects:
        header += [f"{reporting.subject_code(s, subjects)} %" for s in subject_ids]
    header += ["Present", "Total", "Overall %", "Status"]
    ws.append(header)
    for cell in ws[3]:
        cell.font = Font(bold=True)

    red_fill = PatternFill("solid", fgColor="F9D6D3")
    for r in rows:
        line = [r["roll"], r["name"]]
        if show_subjects:
            for sid in subject_ids:
                present, total = r["per_subject"][sid]
                line.append(round(present / total * 100, 1) if total else "-")
        line += [r["present"], r["total"],
                 round(r["pct"], 1) if r["pct"] is not None else "-",
                 reporting.status_text(r["pct"], threshold)]
        ws.append(line)
        if reporting.is_low(r["pct"], threshold):
            for cell in ws[ws.max_row]:
                cell.fill = red_fill
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 26
    wb.save(save_path)


# ---------------------------------------------------------------- picker

class SubjectPicker(tb.Frame):
    """
    Branch -> Semester -> Subject (or "All Subjects") selector shared by every
    report tab, plus the live "Minimum attendance %" box.
    """

    def __init__(self, parent, on_change, on_threshold_change):
        super().__init__(parent, padding=(0, 0, 0, 10))
        self.on_change = on_change
        self.on_threshold_change = on_threshold_change
        self._subject_ids_by_label = {}
        self._threshold = float(config.ATTENDANCE_ELIGIBILITY_THRESHOLD)
        self._pending = None

        tb.Label(self, text="Branch:").pack(side="left")
        self.branch_combo = ttk.Combobox(self, state="readonly", width=8)
        self.branch_combo.pack(side="left", padx=(5, 12))
        self.branch_combo.bind("<<ComboboxSelected>>", lambda e: self._refresh_subjects())

        tb.Label(self, text="Semester:").pack(side="left")
        self.semester_combo = ttk.Combobox(self, state="readonly", width=5, values=config.SEMESTERS)
        self.semester_combo.pack(side="left", padx=(5, 12))
        self.semester_combo.bind("<<ComboboxSelected>>", lambda e: self._refresh_subjects())

        tb.Label(self, text="Subject:").pack(side="left")
        self.subject_combo = ttk.Combobox(self, state="readonly", width=34)
        self.subject_combo.pack(side="left", padx=5)
        self.subject_combo.bind("<<ComboboxSelected>>", lambda e: self.on_change())

        # ---- live minimum-attendance box (right side) ----
        box = tb.Frame(self)
        box.pack(side="right")
        tb.Label(box, text="Minimum attendance %:", font=BOLD).pack(side="left")
        self.threshold_var = tk.StringVar(value=f"{self._threshold:g}")
        self.threshold_spin = ttk.Spinbox(box, from_=0, to=100, increment=5, width=5,
                                          textvariable=self.threshold_var, justify="center")
        self.threshold_spin.pack(side="left", padx=5)
        tb.Button(box, text="Reset", bootstyle="secondary-outline",
                  command=self.reset_threshold).pack(side="left")
        self.threshold_var.trace_add("write", self._on_threshold_typed)

    # ---- threshold ----
    def threshold(self):
        return self._threshold

    def reset_threshold(self):
        self.threshold_var.set(f"{config.ATTENDANCE_ELIGIBILITY_THRESHOLD:g}")

    def _on_threshold_typed(self, *_):
        # wait ~0.3s after the last keystroke so typing "72" doesn't refresh for "7" first
        if self._pending is not None:
            self.after_cancel(self._pending)
        self._pending = self.after(300, self._apply_threshold)

    def _apply_threshold(self):
        self._pending = None
        value = reporting.parse_threshold(self.threshold_var.get(), None)
        if value is None:
            return   # empty / invalid / outside 0-100: keep the last valid minimum
        if value != self._threshold:
            self._threshold = value
            self.on_threshold_change()

    # ---- branch / semester / subject ----
    def refresh(self):
        branches = storage.load_branches()
        self.branch_combo["values"] = branches
        if branches and not self.branch_combo.get():
            self.branch_combo.set(branches[0])
        if not self.semester_combo.get():
            self.semester_combo.set(config.DEFAULT_SEMESTER)
        self._refresh_subjects()

    def _refresh_subjects(self):
        branch, semester = self.branch_combo.get(), self.semester_combo.get()
        subjects = storage.subjects_for(branch, semester) if branch and semester else {}
        # subjects stay in the order they were added (theory first, then labs)
        self._subject_ids_by_label = {info["name"]: sid for sid, info in subjects.items()}
        names = ([ALL_SUBJECTS_LABEL] + list(self._subject_ids_by_label.keys())) if subjects else []
        current = self.subject_combo.get()
        self.subject_combo["values"] = names
        if names and current not in names:
            self.subject_combo.set(names[0])
        elif not names:
            self.subject_combo.set("")
        self.on_change()

    def subject_ids(self):
        """subject_ids the current selection covers (several for All Subjects), in the order they were added."""
        selection = self.subject_combo.get()
        if not selection:
            return []
        if selection == ALL_SUBJECTS_LABEL:
            return list(self._subject_ids_by_label.values())
        sid = self._subject_ids_by_label.get(selection)
        return [sid] if sid else []

    def is_all(self):
        return self.subject_combo.get() == ALL_SUBJECTS_LABEL

    def subject_name(self, subject_id):
        for name, sid in self._subject_ids_by_label.items():
            if sid == subject_id:
                return name
        return subject_id

    def branch_semester(self):
        return self.branch_combo.get(), self.semester_combo.get()

    def heading(self):
        branch, semester = self.branch_semester()
        return f"{branch} Sem {semester} - {self.subject_combo.get() or 'no subject'}"


# ---------------------------------------------------------------- screen

class ReportScreen(tb.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller

        make_header(self, "Attendance Report", controller)

        self.picker = SubjectPicker(self, self._on_subject_change, self._on_threshold_change)
        self.picker.pack(fill="x", padx=15, pady=(12, 0))

        notebook = ttk.Notebook(self)
        notebook.pack(fill="both", expand=True, padx=15, pady=(5, 15))

        self.daily_tab = DailyTab(notebook, self.picker)
        self.weekly_tab = WeeklyTab(notebook, self.picker)
        self.monthly_tab = MonthlyTab(notebook, self.picker)
        self.semester_tab = SemesterTab(notebook, self.picker)

        notebook.add(self.semester_tab, text="Semester Summary")
        notebook.add(self.daily_tab, text="Daily")
        notebook.add(self.weekly_tab, text="Weekly")
        notebook.add(self.monthly_tab, text="Monthly")

        self._tabs = (self.daily_tab, self.weekly_tab, self.monthly_tab, self.semester_tab)

    def _on_subject_change(self):
        for tab in self._tabs:
            tab.on_subject_change()

    def _on_threshold_change(self):
        for tab in self._tabs:
            tab.on_threshold_change()

    def on_show(self):
        self.picker.refresh()

    def on_leave(self):
        pass


# ---------------------------------------------------------------- Daily

class DailyTab(tb.Frame):
    def __init__(self, parent, picker):
        super().__init__(parent, padding=15)
        self.picker = picker

        top = tb.Frame(self)
        top.pack(fill="x", pady=(0, 10))
        tb.Label(top, text="Select date:").pack(side="left", padx=(0, 8))
        self.date_combo = ttk.Combobox(top, width=18, state="readonly")
        self.date_combo.pack(side="left", padx=5)
        self.date_combo.bind("<<ComboboxSelected>>", lambda e: self.load())
        tb.Button(top, text="Export to Excel", bootstyle="success", command=self.export_excel).pack(side="left", padx=5)
        self.summary_label = tb.Label(top, text="", font=("Segoe UI", 11, "bold"))
        self.summary_label.pack(side="right")

        body = tb.Frame(self)
        body.pack(fill="both", expand=True)
        body.columnconfigure(0, weight=1)
        body.columnconfigure(1, weight=1)
        body.rowconfigure(0, weight=1)

        chart_frame = tb.Labelframe(body, text="Present vs Absent", padding=10)
        chart_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 8))
        self.fig = Figure(figsize=(4, 3), dpi=90)
        self.ax = self.fig.add_subplot(111)
        self.canvas = FigureCanvasTkAgg(self.fig, master=chart_frame)
        self.canvas.get_tk_widget().pack(fill="both", expand=True)

        table_frame = tb.Labelframe(body, text="Details", padding=10)
        table_frame.grid(row=0, column=1, sticky="nsew", padx=(8, 0))
        table, self.tree = make_table(
            table_frame, ("subject", "roll", "name", "time", "status"),
            ("Subject", "Roll No", "Name", "Time", "Status"),
            {"subject": 130, "roll": 70, "name": 150, "time": 80, "status": 80},
        )
        table.pack(fill="both", expand=True)

    def on_threshold_change(self):
        pass   # daily view has no percentages

    def on_subject_change(self):
        subject_ids = self.picker.subject_ids()
        dates = storage.list_dates_for_subjects(subject_ids) if subject_ids else []
        self.date_combo["values"] = dates
        if dates and self.date_combo.get() not in dates:
            self.date_combo.set(dates[-1])
        elif not dates:
            self.date_combo.set("")
        self.load()

    def load(self):
        subject_ids = self.picker.subject_ids()
        date = self.date_combo.get()
        _clear_tree(self.tree)
        if not subject_ids:
            self.summary_label.config(text="")
            _empty_chart(self.fig, self.ax, self.canvas, "No subject selected - add one in Manage Subjects")
            return
        if not date:
            self.summary_label.config(text="")
            _empty_chart(self.fig, self.ax, self.canvas, "No attendance recorded yet")
            return

        present = absent = 0
        for subject_id in subject_ids:
            subject_name = self.picker.subject_name(subject_id)
            for roll_no, name, time_str, status in [r[:4] for r in storage.read_subject_rows(date, subject_id)]:
                tag = "present" if status == "Present" else "absent"
                self.tree.insert("", "end", values=(subject_name, roll_no, name, time_str, status), tags=(tag,))
                if status == "Present":
                    present += 1
                else:
                    absent += 1

        self.summary_label.config(text=f"Present: {present}   Absent: {absent}")
        self.ax.clear()
        self.ax.bar(["Present", "Absent"], [present, absent], color=[GREEN, RED])
        self.ax.set_ylabel("Students")
        self.fig.tight_layout()
        self.canvas.draw()

    def export_excel(self):
        subject_ids = self.picker.subject_ids()
        date = self.date_combo.get()
        if not subject_ids or not date:
            messagebox.showinfo("No data", "Pick a subject and a date first.")
            return
        save_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", initialfile=f"attendance_{date}.xlsx",
            filetypes=[("Excel files", "*.xlsx")])
        if not save_path:
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = date
        ws.append(["Subject", "Roll No", "Name", "Time", "Status"])
        for subject_id in subject_ids:
            subject_name = self.picker.subject_name(subject_id)
            for row in storage.read_subject_rows(date, subject_id):
                ws.append([subject_name] + list(row[:4]))
        wb.save(save_path)
        messagebox.showinfo("Exported", f"Report saved to:\n{save_path}")


# ---------------------------------------------------------------- Weekly / Monthly

class _PeriodTab(tb.Frame):
    """Shared body of the Weekly and Monthly tabs - they differ only in how dates are grouped and charted."""

    PERIOD_LABEL = ""      # "week" / "month"
    CHART_TITLE = ""

    def __init__(self, parent, picker):
        super().__init__(parent, padding=15)
        self.picker = picker
        self._rows = []
        self._period_dates = []
        self._key_by_label = {}

        top = tb.Frame(self)
        top.pack(fill="x", pady=(0, 10))
        tb.Label(top, text=f"Select {self.PERIOD_LABEL}:").pack(side="left", padx=(0, 8))
        self.period_combo = ttk.Combobox(top, width=24, state="readonly")
        self.period_combo.pack(side="left", padx=5)
        self.period_combo.bind("<<ComboboxSelected>>", lambda e: self.load())
        tb.Button(top, text="Export to Excel", bootstyle="success", command=self.export_excel).pack(side="left", padx=5)
        self.legend_label = tb.Label(top, text="", bootstyle="secondary")
        self.legend_label.pack(side="right")

        body = tb.Frame(self)
        body.pack(fill="both", expand=True)
        body.rowconfigure(1, weight=1)
        body.columnconfigure(0, weight=1)

        chart_frame = tb.Labelframe(body, text=self.CHART_TITLE, padding=10)
        chart_frame.grid(row=0, column=0, sticky="ew", pady=(0, 10))
        self.fig = Figure(figsize=(8, 2.6), dpi=90)
        self.ax = self.fig.add_subplot(111)
        self.canvas = FigureCanvasTkAgg(self.fig, master=chart_frame)
        self.canvas.get_tk_widget().pack(fill="both", expand=True)

        table_frame = tb.Labelframe(body, text="Per-student summary", padding=10)
        table_frame.grid(row=1, column=0, sticky="nsew")
        table, self.tree = make_table(
            table_frame, ("roll", "name", "present", "total", "pct", "status"),
            ("Roll No", "Name", "Present", "Total", "%", "Status"),
            {"roll": 90, "name": 200, "present": 80, "total": 80, "pct": 80, "status": 80},
            numeric_columns=("roll", "present", "total", "pct"),
        )
        table.pack(fill="both", expand=True)
        self.tree.column("name", anchor="w")

    # -- subclass hooks --
    def _key(self, date_str):
        raise NotImplementedError

    def _label(self, key):
        return key

    def _draw_chart(self, period_dates, subject_ids):
        raise NotImplementedError

    def _selected_key(self):
        return self._key_by_label.get(self.period_combo.get(), "")

    # -- common flow --
    def on_subject_change(self):
        subject_ids = self.picker.subject_ids()
        all_dates = storage.list_dates_for_subjects(subject_ids) if subject_ids else []
        keys = sorted({self._key(d) for d in all_dates}, reverse=True)
        self._key_by_label = {self._label(k): k for k in keys}
        periods = [self._label(k) for k in keys]
        current = self.period_combo.get()
        self.period_combo["values"] = periods
        if periods and current not in periods:
            self.period_combo.set(periods[0])
        elif not periods:
            self.period_combo.set("")
        self.load()

    def on_threshold_change(self):
        self._render_table()

    def load(self):
        subject_ids = self.picker.subject_ids()
        period = self._selected_key()
        self._rows, self._period_dates = [], []
        _clear_tree(self.tree)
        if not subject_ids:
            _empty_chart(self.fig, self.ax, self.canvas, "No subject selected - add one in Manage Subjects")
            return
        if not period:
            _empty_chart(self.fig, self.ax, self.canvas, "No attendance recorded yet")
            return

        all_dates = storage.list_dates_for_subjects(subject_ids)
        self._period_dates = sorted(d for d in all_dates if self._key(d) == period)
        self.ax.clear()
        self._draw_chart(self._period_dates, subject_ids)

        branch, semester = self.picker.branch_semester()
        eligible = storage.students_for_subjects(branch, semester, subject_ids)
        self._rows = reporting.student_matrix(self._period_dates, subject_ids, eligible)
        self._render_table()

    def _render_table(self):
        threshold = self.picker.threshold()
        self.legend_label.config(text=f"Red = below {threshold:g}%")
        _fill_simple_tree(self.tree, self._rows, threshold)

    def export_excel(self):
        subject_ids = self.picker.subject_ids()
        period = self._selected_key()
        if not subject_ids or not period or not self._rows:
            messagebox.showinfo("No data", f"Pick a subject and a {self.PERIOD_LABEL} first.")
            return
        save_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", initialfile=f"attendance_{self.PERIOD_LABEL}_{period}.xlsx",
            filetypes=[("Excel files", "*.xlsx")])
        if not save_path:
            return
        _save_report_excel(save_path, period, self._rows, subject_ids, self.picker.threshold(),
                           f"{self.picker.heading()} - {self.period_combo.get()}")
        messagebox.showinfo("Exported", f"Report saved to:\n{save_path}")


class WeeklyTab(_PeriodTab):
    PERIOD_LABEL = "week"
    CHART_TITLE = "Attendance by day"

    def _key(self, date_str):
        return _week_key(date_str)

    def _label(self, key):
        return _week_label(key)

    def _draw_chart(self, period_dates, subject_ids):
        counts = storage.daily_counts(period_dates, subject_ids)
        labels = [datetime.strptime(d, "%Y-%m-%d").strftime("%a %d") for d, _, _ in counts]
        present_vals = [p for _, p, _ in counts]
        absent_vals = [a for _, _, a in counts]
        x = list(range(len(labels)))
        width = 0.35
        self.ax.bar([i - width / 2 for i in x], present_vals, width, label="Present", color=GREEN)
        self.ax.bar([i + width / 2 for i in x], absent_vals, width, label="Absent", color=RED)
        self.ax.set_xticks(x)
        self.ax.set_xticklabels(labels)
        if labels:
            self.ax.legend()
        self.fig.tight_layout()
        self.canvas.draw()


class MonthlyTab(_PeriodTab):
    PERIOD_LABEL = "month"
    CHART_TITLE = "Daily present-count trend"

    def _key(self, date_str):
        return _month_key(date_str)

    def _label(self, key):
        return _month_label(key)

    def _draw_chart(self, period_dates, subject_ids):
        counts = storage.daily_counts(period_dates, subject_ids)
        days = [datetime.strptime(d, "%Y-%m-%d").day for d, _, _ in counts]
        present_vals = [p for _, p, _ in counts]
        self.ax.plot(days, present_vals, marker="o", color=GREEN)
        self.ax.set_xlabel("Day of month")
        self.ax.set_ylabel("Present")
        self.fig.tight_layout()
        self.canvas.draw()


# ---------------------------------------------------------------- Semester Summary

class SemesterTab(tb.Frame):
    """
    Whole-semester view. By default it covers EVERY day on which attendance was taken
    (change From / To to look at a smaller period). Nothing to click - it updates by itself
    when the subject, the dates or the minimum attendance % change.
    """

    def __init__(self, parent, picker):
        super().__init__(parent, padding=15)
        self.picker = picker
        self._rows = []
        self._subject_ids = []
        self._range_dates = []
        self._table_frame = None
        self._table_sig = None
        self.tree = None

        # ---- row 1: period + actions ----
        top = tb.Frame(self)
        top.pack(fill="x")
        tb.Label(top, text="Period:  From").pack(side="left")
        self.from_combo = ttk.Combobox(top, width=12, state="readonly")
        self.from_combo.pack(side="left", padx=(5, 10))
        self.from_combo.bind("<<ComboboxSelected>>", lambda e: self.load())
        tb.Label(top, text="To").pack(side="left")
        self.to_combo = ttk.Combobox(top, width=12, state="readonly")
        self.to_combo.pack(side="left", padx=(5, 10))
        self.to_combo.bind("<<ComboboxSelected>>", lambda e: self.load())
        tb.Button(top, text="Full semester", bootstyle="secondary-outline",
                  command=self.reset_range).pack(side="left", padx=5)
        tb.Button(top, text="Export to Excel", bootstyle="success",
                  command=self.export_excel).pack(side="left", padx=5)
        self.only_low_var = tk.BooleanVar(value=False)
        tb.Checkbutton(top, text="Show only students below minimum", variable=self.only_low_var,
                       command=self._render, bootstyle="danger-round-toggle").pack(side="right")

        self.info_label = tb.Label(self, text="", bootstyle="secondary")
        self.info_label.pack(anchor="w", pady=(6, 8))

        # ---- row 2: summary cards ----
        cards = tb.Frame(self)
        cards.pack(fill="x")
        for i in range(4):
            cards.columnconfigure(i, weight=1, uniform="card")
        self.students_var = tk.StringVar(value="0")
        self.ok_var = tk.StringVar(value="0")
        self.low_var = tk.StringVar(value="0")
        self.avg_var = tk.StringVar(value="-")
        self.ok_title = tk.StringVar(value="Meeting minimum")
        self.low_title = tk.StringVar(value="Below minimum")
        self._card(cards, 0, self.students_var, "Students", "info")
        self._card(cards, 1, self.ok_var, self.ok_title, "success")
        self._card(cards, 2, self.low_var, self.low_title, "danger")
        self._card(cards, 3, self.avg_var, "Class average", "secondary")

        # ---- row 3: charts ----
        chart_frame = tb.Labelframe(self, text="Overview", padding=8)
        chart_frame.pack(fill="x", pady=(10, 8))
        self.fig = Figure(figsize=(8, 2.5), dpi=90)
        self.canvas = FigureCanvasTkAgg(self.fig, master=chart_frame)
        self.canvas.get_tk_widget().pack(fill="both", expand=True)

        # ---- row 4: student table (rebuilt when the subject columns change) ----
        self.table_holder = tb.Labelframe(self, text="Student-wise attendance  (▼ = below minimum)", padding=8)
        self.table_holder.pack(fill="both", expand=True)

    def _card(self, parent, col, var, title, style):
        card = tb.Frame(parent, bootstyle=style, padding=(10, 8))
        card.grid(row=0, column=col, padx=5, sticky="ew")
        tb.Label(card, textvariable=var, font=("Segoe UI", 20, "bold"), bootstyle=f"inverse-{style}").pack()
        if isinstance(title, tk.StringVar):
            tb.Label(card, textvariable=title, font=("Segoe UI", 10), bootstyle=f"inverse-{style}").pack()
        else:
            tb.Label(card, text=title, font=("Segoe UI", 10), bootstyle=f"inverse-{style}").pack()

    # ---- picker / threshold hooks ----
    def on_subject_change(self):
        subject_ids = self.picker.subject_ids()
        dates = storage.list_dates_for_subjects(subject_ids) if subject_ids else []
        self.from_combo["values"] = dates
        self.to_combo["values"] = dates
        if dates:
            if self.from_combo.get() not in dates:
                self.from_combo.set(dates[0])
            if self.to_combo.get() not in dates:
                self.to_combo.set(dates[-1])
        else:
            self.from_combo.set("")
            self.to_combo.set("")
        self.load()

    def on_threshold_change(self):
        self._render()

    def reset_range(self):
        dates = list(self.from_combo["values"])
        if dates:
            self.from_combo.set(dates[0])
            self.to_combo.set(dates[-1])
        self.load()

    # ---- table ----
    def _ensure_table(self, subject_ids):
        show_subjects = len(subject_ids) > 1
        signature = tuple(subject_ids) if show_subjects else ()
        if self._table_frame is not None and signature == self._table_sig:
            return
        if self._table_frame is not None:
            self._table_frame.destroy()
        subjects = storage.load_subjects()
        columns = ["roll", "name"] + [f"s{i}" for i in range(len(signature))] + ["attended", "pct", "status"]
        headings = (["Roll No", "Name"] + [reporting.subject_code(s, subjects) for s in signature]
                    + ["Attended", "Overall %", "Status"])
        widths = {"roll": 70, "name": 150, "attended": 70, "pct": 70, "status": 60}
        widths.update({f"s{i}": 58 for i in range(len(signature))})
        numeric = tuple(c for c in columns if c not in ("name", "status"))
        self._table_frame, self.tree = make_table(
            self.table_holder, tuple(columns), tuple(headings), widths, numeric_columns=numeric)
        self._table_frame.pack(fill="both", expand=True)
        self.tree.column("name", anchor="w")
        self._table_sig = signature

    # ---- load / render ----
    def load(self):
        subject_ids = self.picker.subject_ids()
        start, end = self.from_combo.get(), self.to_combo.get()
        self._subject_ids = subject_ids
        self._rows, self._range_dates = [], []

        if subject_ids and start and end:
            if start > end:
                start, end = end, start
            all_dates = storage.list_dates_for_subjects(subject_ids)
            self._range_dates = [d for d in all_dates if start <= d <= end]
            branch, semester = self.picker.branch_semester()
            eligible = storage.students_for_subjects(branch, semester, subject_ids)
            self._rows = reporting.student_matrix(self._range_dates, subject_ids, eligible)
        self._render()

    def _render(self):
        threshold = self.picker.threshold()
        rows, subject_ids = self._rows, self._subject_ids
        self._ensure_table(subject_ids)
        _clear_tree(self.tree)

        self.ok_title.set(f"Meeting {threshold:g}%+")
        self.low_title.set(f"Below {threshold:g}%")

        if not subject_ids:
            self.info_label.config(text="No subject selected - add subjects in Manage Subjects.")
        elif not self._range_dates:
            self.info_label.config(text="No attendance recorded yet for this selection.")
        else:
            self.info_label.config(text=(
                f"{self.picker.heading()}  |  {len(self._range_dates)} class days "
                f"({self._range_dates[0]} to {self._range_dates[-1]}).  "
                f"Each subject column = that student's attendance % in that subject; "
                f"Overall = all lectures together.  Red = below {threshold:g}%."))

        stats = reporting.summary_stats(rows, threshold)
        self.students_var.set(str(stats["students"]))
        self.ok_var.set(str(stats["ok"]))
        self.low_var.set(str(stats["low"]))
        self.avg_var.set("-" if stats["average"] is None else f"{stats['average']:.0f}%")

        only_low = self.only_low_var.get()
        subjects_shown = len(subject_ids) > 1
        for r in rows:
            low = reporting.is_low(r["pct"], threshold)
            if only_low and not low:
                continue
            values = [r["roll"], r["name"]]
            if subjects_shown:
                values += [_subject_cell(*r["per_subject"][sid], threshold) for sid in subject_ids]
            values += [f"{r['present']}/{r['total']}", _fmt_pct(r["pct"]),
                       reporting.status_text(r["pct"], threshold)]
            self.tree.insert("", "end", values=tuple(values), tags=("low",) if low else ())

        self._draw(rows, subject_ids, threshold, stats)

    def _draw(self, rows, subject_ids, threshold, stats):
        self.fig.clear()
        if not rows or not stats["with_data"]:
            ax = self.fig.add_subplot(111)
            message = ("No subject selected - add one in Manage Subjects" if not subject_ids else
                       "No students registered for this branch + semester" if not rows else
                       "No attendance recorded yet")
            _empty_chart(self.fig, ax, self.canvas, message)
            return

        # left: how many students are fine vs short
        ax1 = self.fig.add_subplot(121)
        ax1.bar(["Meeting\nminimum", "Below\nminimum"], [stats["ok"], stats["low"]], color=[GREEN, RED])
        for i, value in enumerate([stats["ok"], stats["low"]]):
            ax1.text(i, value, str(value), ha="center", va="bottom", fontsize=9)
        ax1.set_title(f"Students vs minimum ({threshold:g}%)", fontsize=10)
        ax1.set_ylabel("Students")

        # right: class average per subject (All Subjects) or spread of % (single subject)
        ax2 = self.fig.add_subplot(122)
        if len(subject_ids) > 1:
            subjects = storage.load_subjects()
            averages = reporting.subject_averages(rows, subject_ids)
            labels = [reporting.subject_code(s, subjects) for s in subject_ids]
            values = [averages[s] or 0 for s in subject_ids]
            ax2.bar(labels, values, color=[GREEN if v >= threshold else RED for v in values])
            ax2.axhline(threshold, color="#555555", linestyle="--", linewidth=1)
            ax2.set_ylim(0, 100)
            ax2.set_title("Class average % by subject", fontsize=10)
            ax2.tick_params(axis="x", labelrotation=45, labelsize=7)
        else:
            bands = reporting.histogram(rows, 10)
            ax2.bar([f"{lo}-{hi}" for lo, hi, _ in bands], [c for _, _, c in bands],
                    color=[GREEN if (lo + hi) / 2 >= threshold else RED for lo, hi, _ in bands])
            ax2.set_title("Students by attendance % band", fontsize=10)
            ax2.tick_params(axis="x", labelrotation=45, labelsize=7)
        ax2.set_ylabel("%" if len(subject_ids) > 1 else "Students")
        self.fig.tight_layout()
        self.canvas.draw()

    def export_excel(self):
        if not self._subject_ids or not self._rows or not self._range_dates:
            messagebox.showinfo("No data", "Pick a subject with recorded attendance first.")
            return
        save_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", initialfile="attendance_semester_summary.xlsx",
            filetypes=[("Excel files", "*.xlsx")])
        if not save_path:
            return
        _save_report_excel(
            save_path, "Semester Summary", self._rows, self._subject_ids, self.picker.threshold(),
            f"{self.picker.heading()} - {self._range_dates[0]} to {self._range_dates[-1]}")
        messagebox.showinfo("Exported", f"Report saved to:\n{save_path}")
